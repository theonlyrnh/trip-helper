from __future__ import annotations

from decimal import Decimal
from pathlib import Path
import sys

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

ROOT = Path(__file__).resolve().parents[1]
BACKEND = ROOT / "backend"
if str(BACKEND) not in sys.path:
    sys.path.insert(0, str(BACKEND))

from database import Base  # noqa: E402
from models.trip import Trip  # noqa: E402
from models.document import Document  # noqa: E402
from models.invoice import Invoice  # noqa: E402
from enums import ExpenseCategory, ReimbursementStatus  # noqa: E402
from exporters.excel_exporter import ExcelExporter  # noqa: E402
from services.expense_summary import ExpenseSummary  # noqa: E402


@pytest.fixture()
def db_session(tmp_path, monkeypatch):
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    Base.metadata.create_all(bind=engine)
    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    db = SessionLocal()
    monkeypatch.setattr("exporters.excel_exporter.settings.EXPORT_DIR", str(tmp_path))
    try:
        yield db
    finally:
        db.close()
        Base.metadata.drop_all(bind=engine)


def test_excel_export_reimbursement_total_matches_summary_and_contains_status_columns(db_session):
    trip = Trip(title="含非法字符/测试:项目", folder_path="D:/tmp", folder_name="tmp", project_type="DAILY")
    db_session.add(trip)
    db_session.commit()
    db_session.refresh(trip)

    doc1 = Document(trip_id=trip.id, file_name="a.pdf", file_path="D:/tmp/a.pdf", file_ext=".pdf")
    doc2 = Document(trip_id=trip.id, file_name="b.pdf", file_path="D:/tmp/b.pdf", file_ext=".pdf")
    db_session.add_all([doc1, doc2])
    db_session.commit()
    db_session.refresh(doc1)
    db_session.refresh(doc2)

    db_session.add_all([
        Invoice(
            trip_id=trip.id,
            document_id=doc1.id,
            invoice_type="GENERAL_INVOICE",
            expense_category=ExpenseCategory.MEAL,
            total_amount=Decimal("100.00"),
            confirmed_amount=Decimal("88.00"),
            reimbursement_status=ReimbursementStatus.THIS_TRIP,
            include_in_summary=True,
            document_role="OFFICIAL_INVOICE",
            confidence=1.0,
        ),
        Invoice(
            trip_id=trip.id,
            document_id=doc2.id,
            invoice_type="GENERAL_INVOICE",
            expense_category=ExpenseCategory.MEAL,
            total_amount=Decimal("200.00"),
            reimbursement_status=ReimbursementStatus.ALREADY_REIMBURSED,
            include_in_summary=False,
            document_role="OFFICIAL_INVOICE",
            confidence=1.0,
        ),
    ])
    db_session.commit()
    ExpenseSummary.update_trip(db_session, trip)

    path = ExcelExporter.export(db_session, trip)

    assert Path(path).exists()
    assert ":" not in Path(path).name
    workbook = load_workbook(path, data_only=True)
    summary_sheet = workbook["项目摘要"]
    assert summary_sheet["B15"].value == 88.0
    detail_sheet = workbook["发票明细"]
    headers = [cell.value for cell in detail_sheet[1]]
    assert "是否计入本次报销" in headers
    assert "报销状态" in headers
    assert "本次报销金额" in headers
    included_amount_col = headers.index("本次报销金额") + 1
    exported_total = sum(
        row[included_amount_col - 1]
        for row in detail_sheet.iter_rows(min_row=2, values_only=True)
        if isinstance(row[included_amount_col - 1], (int, float))
    )
    assert exported_total == 88.0
